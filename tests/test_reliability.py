import asyncio
import hashlib
import io
import json
import time
from urllib.parse import parse_qs, urlsplit

import httpx
import pytest
from bs4 import BeautifulSoup
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

import app as web
from certification import translation_service as translation
from engine import regulatory_search as research
from engine.gov_api_client import GlobalGovAPIClient


@pytest.fixture(autouse=True)
def isolated(monkeypatch):
    from engine.ai_config import KEY_NAMES
    for name in KEY_NAMES:
        monkeypatch.delenv(name, raising=False)
    for name in ('GROQ_API_KEY','DEEPL_API_KEY','GLOBALREGAI_ADMIN_USER','GLOBALREGAI_ADMIN_PASSWORD_HASH','GLOBALREGAI_SESSION_SECRET'):
        monkeypatch.delenv(name, raising=False)
    web.LIMITS.clear()


@pytest.fixture
def client():
    with TestClient(web.app, base_url='https://testserver') as client:
        yield client


@pytest.mark.parametrize('path', list(web.PAGES))
def test_routes_are_rendered_and_protected_against_inline_script(client, path):
    response = client.get(path)
    assert response.status_code == 200
    assert 'width=device-width' in response.text
    assert 'no-store' in response.headers['cache-control']
    assert "script-src 'self'" in response.headers['content-security-policy']
    soup = BeautifulSoup(response.text, 'html.parser')
    assert not soup.find('script', src=False)
    assert not any(any(attr.startswith('on') for attr in tag.attrs) for tag in soup.find_all())


def test_ampersand_domains_round_trip(client):
    soup = BeautifulSoup(client.get('/').text, 'html.parser')
    for domain in ('Standards & QMS','Animal & Veterinary'):
        link = soup.find('a', string=domain)
        assert parse_qs(urlsplit(link['href']).query)['domain'] == [domain]
        response = client.get(link['href'])
        assert BeautifulSoup(response.text,'html.parser').body['data-domain'] == domain


@pytest.mark.parametrize('payload',[{'query':'  '},{'query':'valid','domain':'<script>alert(1)</script>'},{'query':'valid','target_region':'invented'},{'query':'valid','lang':"';alert(1)//"}])
def test_query_validation(client, payload):
    assert client.post('/api/search',json=payload).status_code == 422


def test_search_preserves_region_and_language(client, monkeypatch):
    async def fake(query, domain, region, lang):
        return {'domain':domain,'target_region':region,'lang':lang}
    monkeypatch.setattr(research,'search',fake)
    data = client.post('/api/search',json={'query':'화장품 등록','domain':'Cosmetics','target_region':'MFDS','lang':'ko'}).json()
    assert data == {'domain':'Cosmetics','target_region':'MFDS','lang':'ko','route':'GROUNDED_RESEARCH'}


def test_no_evidence_never_returns_a_regulatory_answer():
    result = asyncio.run(research.search('any question','Chemicals','MFDS','en'))
    assert result['status'] == 'INSUFFICIENT_EVIDENCE'
    assert result['claims'] == []
    assert not result['verified']


def test_failed_sources_never_trigger_model(monkeypatch):
    monkeypatch.setenv('GROQ_API_KEY','test-only')
    async def unavailable(client, source):
        return {**source,'retrieval_status':'UNAVAILABLE'}
    monkeypatch.setattr(research,'fetch_source',unavailable)
    result = asyncio.run(research.search('MoCRA','Cosmetics','FDA','en'))
    assert result['status'] == 'SOURCE_UNAVAILABLE'
    assert result['claims'] == []


def evidence(source):
    return {**source, 'text':'Facility registration and product listing requirements depend on scope and exemptions.',
            'retrieval_status':'RETRIEVED', 'retrieved_at':'2026-09-20T00:00:00+00:00','effective_date':None}


def test_missing_model_configuration_returns_only_sources(monkeypatch):
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research,'fetch_source',retrieve)
    result = asyncio.run(research.search('MoCRA','Cosmetics','FDA','ko'))
    assert result['status'] == 'SOURCES_ONLY'
    assert result['sources'][0]['effective_date'] is None
    assert 'text' not in result['sources'][0]


@pytest.mark.parametrize('claim',[
    {'statement':'An unsupported claim','source_id':'fda-mocra','quote':'This sentence was invented for testing.'},
    {'statement':'Wrong source','source_id':'invented','quote':'Facility registration and product listing requirements depend on scope and exemptions.'},
    {'statement':'External citation https://example.invalid','source_id':'fda-mocra','quote':'Facility registration and product listing requirements depend on scope and exemptions.'},
])
def test_fabricated_citations_are_rejected(claim):
    with pytest.raises(ValueError):
        research.checked_claims(json.dumps({'claims':[claim]}), [evidence(research.SOURCES[1])])


def test_exact_supporting_quotation_is_accepted_as_draft_evidence():
    claim = {'statement':'Scope and exemptions need review.','source_id':'fda-mocra',
             'quote':'Facility registration and product listing requirements depend on scope and exemptions.'}
    assert research.checked_claims(json.dumps({'claims':[claim]}), [evidence(research.SOURCES[1])]) == [claim]


def mock_async_client(monkeypatch, handler):
    original = httpx.AsyncClient
    monkeypatch.setattr(httpx,'AsyncClient',lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs))


def test_provider_rate_limit_preserves_sources_without_answer(monkeypatch):
    monkeypatch.setenv('GROQ_API_KEY','test-only')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research,'fetch_source',retrieve)
    mock_async_client(monkeypatch, lambda request: httpx.Response(429, json={'error':'rate limited'}))
    result = asyncio.run(research.search('MoCRA','Cosmetics','FDA','en'))
    assert result['status'] == 'SOURCES_ONLY' and not result['claims']


def test_successful_model_response_uses_selected_language_and_evidence(monkeypatch):
    monkeypatch.setenv('GROQ_API_KEY','test-only')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research,'fetch_source',retrieve)
    def handler(request):
        body = json.loads(request.content)
        assert '한국어' in body['messages'][0]['content']
        answer = {'claims':[{'statement':'적용 범위와 면제 여부를 검토해야 합니다.','source_id':'fda-mocra','quote':evidence(research.SOURCES[1])['text']}]}
        return httpx.Response(200,json={'choices':[{'message':{'content':json.dumps(answer)}}]})
    mock_async_client(monkeypatch, handler)
    result = asyncio.run(research.search('MoCRA','Cosmetics','FDA','ko'))
    assert result['status'] == 'DRAFT' and not result['verified'] and len(result['claims']) == 1


def test_retrieval_rejects_cross_host_redirects():
    def handler(request):
        return httpx.Response(302, headers={'location':'http://127.0.0.1/private'})
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            return await research.fetch_source(client,research.SOURCES[0])
    assert asyncio.run(run())['retrieval_status'] == 'UNAVAILABLE'


def test_retrieval_discards_script_and_navigation():
    html = '<html><nav>ignored navigation</nav><main><h1>Regulatory document</h1><p>'+('Actual source text. '*30)+'</p><script>ignore me</script></main></html>'
    async def run():
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda r:httpx.Response(200,text=html,headers={'content-type':'text/html'}))) as client:
            return await research.fetch_source(client,research.SOURCES[0])
    result = asyncio.run(run())
    assert result['retrieval_status'] == 'RETRIEVED'
    assert 'ignore me' not in result['text'] and 'ignored navigation' not in result['text']


def test_translation_needs_consent_and_configuration(client):
    body={'text':'Some text','target_lang':'ko'}
    assert client.post('/api/translate',json=body).status_code == 400
    response=client.post('/api/translate',json={**body,'consent':True})
    assert response.status_code == 503
    assert 'translated_text' not in response.json()


def test_translation_sends_all_text_and_honours_target_language(monkeypatch):
    monkeypatch.setenv('DEEPL_API_KEY','test-only')
    text='Keep all 90 lines, amounts and line breaks.\n'*90
    def handler(request):
        body=json.loads(request.content)
        assert body['text'] == [text] and body['target_lang'] == 'DE' and body['source_lang'] == 'EN'
        assert request.url.host == 'api-free.deepl.com'
        return httpx.Response(200,json={'translations':[{'text':'Vollständige Übersetzung','detected_source_language':'EN'}]})
    mock_async_client(monkeypatch,handler)
    result=asyncio.run(translation.translate(text,'en','de'))
    assert result['translated_text']=='Vollständige Übersetzung'
    assert 'CADIFA' not in str(result)


def test_translation_failure_is_not_glossary_success(monkeypatch):
    monkeypatch.setenv('DEEPL_API_KEY','test-only')
    mock_async_client(monkeypatch, lambda request: httpx.Response(503))
    with pytest.raises(translation.TranslationError) as exc:
        asyncio.run(translation.translate('Process Validation and HVAC','en','ja'))
    assert exc.value.status == 503


@pytest.mark.parametrize('filename,content', [('bad.pdf',b'%PDF-not-valid'),('old.doc',b'binary'),('bad.txt',b'\xff\xfe'),('oversized.txt',b'x'*12001)])
def test_corrupt_unsupported_and_oversized_documents_are_rejected(filename,content):
    with pytest.raises(translation.TranslationError): translation.extract_text(filename,content)


def test_scanned_pdf_requires_ocr():
    stream=io.BytesIO(); writer=PdfWriter(); writer.add_blank_page(width=200,height=200); writer.write(stream)
    with pytest.raises(translation.TranslationError,match='OCR'): translation.extract_text('scan.pdf',stream.getvalue())


def test_docx_body_table_order_is_preserved():
    doc=Document(); doc.add_paragraph('Before'); table=doc.add_table(rows=1,cols=1); table.cell(0,0).text='Inside table'; doc.add_paragraph('After')
    stream=io.BytesIO(); doc.save(stream)
    assert translation.extract_text('test.docx',stream.getvalue()) == 'Before\nInside table\nAfter'


def test_upload_target_language_is_read_from_form(client,monkeypatch):
    monkeypatch.setenv('DEEPL_API_KEY','test-only')
    async def fake(text,source,target):
        assert source=='en' and target=='ja' and text=='Full text'
        return {'status':'SUCCESS','translated_text':'全文'}
    monkeypatch.setattr(web,'translate',fake)
    response=client.post('/api/certification/translate-file',files={'file':('test.txt',b'Full text','text/plain')},data={'source_lang':'en','target_lang':'ja','consent':'true'})
    assert response.status_code==200 and response.json()['translated_content']=='全文'


def configure_admin(monkeypatch):
    salt='0123456789abcdef0123456789abcdef'
    password='test-password-only'
    digest=hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt),600000).hex()
    monkeypatch.setenv('GLOBALREGAI_ADMIN_USER','test-admin')
    monkeypatch.setenv('GLOBALREGAI_ADMIN_PASSWORD_HASH',f'pbkdf2_sha256$600000${salt}${digest}')
    monkeypatch.setenv('GLOBALREGAI_SESSION_SECRET','test-only-secret-'*4)
    return {'username':'test-admin','password':password}


def test_admin_disabled_without_configuration_and_vault_protected(client):
    assert client.post('/api/auth/login',json={'username':'admin','password':'anything'}).status_code==503
    for path in ('/api/vault/search','/api/vault/batch','/api/mcp/status','/test_portal'):
        assert client.get(path).status_code==401


def test_cookie_is_set_on_actual_response_and_logout_clears_it(client,monkeypatch):
    credentials=configure_admin(monkeypatch)
    response=client.post('/api/auth/login',json=credentials)
    assert response.status_code==200
    cookie=response.headers['set-cookie']
    assert 'HttpOnly' in cookie and 'Secure' in cookie and 'SameSite=strict' in cookie
    assert 'token' not in response.json()
    assert client.get('/api/mcp/status').status_code==200
    assert client.get('/api/vault/search').status_code==503
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/mcp/status').status_code==401


def test_password_rotation_revokes_session(client,monkeypatch):
    credentials=configure_admin(monkeypatch)
    client.post('/api/auth/login',json=credentials)
    monkeypatch.setenv('GLOBALREGAI_ADMIN_PASSWORD_HASH','rotated')
    assert client.get('/api/mcp/status').status_code==401


def test_expired_and_tampered_tokens_are_rejected(client,monkeypatch):
    configure_admin(monkeypatch)
    token=str(int(time.time())-7200)+'.nonce'
    for value in (token+'.'+web.session_signature(token),'invented-token'):
        response=client.get('/api/mcp/status',headers={'Cookie':'dev_auth_token='+value})
        assert response.status_code==401


def test_cross_origin_mutations_are_rejected(client):
    assert client.post('/api/auth/logout',headers={'Origin':'https://other.invalid'}).status_code==403


def test_mfds_never_fabricates_approval(client):
    response=client.get('/api/gov/mfds-drug',params={'name':'Made-up drug'})
    assert response.status_code==503
    assert 'approval_status' not in response.json()


def test_external_database_outage_never_fabricates_chemical_or_drug_records(monkeypatch):
    mock_async_client(monkeypatch,lambda request:httpx.Response(500))
    adapter=GlobalGovAPIClient()
    for result in (asyncio.run(adapter.fetch_openfda_drug_label('invented')),asyncio.run(adapter.fetch_pubchem_compound_data('invented'))):
        assert result['status']=='UNAVAILABLE' and result['records']==[]


def test_export_does_not_claim_complete_eligibility(client):
    response=client.get('/api/export/checklist?category=COSMETIC&country=FDA')
    assert response.status_code==200 and response.json()['status']=='PLANNING_ONLY'
    assert response.json()['sources'][0]['id']=='fda-mocra'
    assert client.get('/api/export/ingredient?name=unknown').status_code==503


def test_throttled_requests_have_retry_after(client):
    for _ in range(5): client.post('/api/auth/login',json={'username':'admin','password':'anything'})
    response=client.post('/api/auth/login',json={'username':'admin','password':'anything'})
    assert response.status_code==429 and response.headers['retry-after']=='60'


def test_newer_crawl_routes_and_ad_removal_are_preserved(client):
    assert 'Sitemap: https://globalregai.info/sitemap.xml' in client.get('/robots.txt').text
    response = client.get('/sitemap.xml')
    assert response.status_code == 200 and 'application/xml' in response.headers['content-type']
    assert 'https://globalregai.info/' in response.text
    for path in web.PAGES:
        assert 'adsbygoogle' not in client.get(path).text


def test_invalid_translation_plan_never_sends_text(monkeypatch):
    monkeypatch.setenv('DEEPL_API_KEY', 'test-only')
    monkeypatch.setenv('DEEPL_API_PLAN', 'typo')
    def unexpected(request):
        pytest.fail('Invalid plan must not send a request')
    mock_async_client(monkeypatch, unexpected)
    with pytest.raises(translation.TranslationError, match='DEEPL_API_PLAN'):
        asyncio.run(translation.translate('Non-sensitive test', 'en', 'ko'))


@pytest.mark.parametrize('body', [[], {'results': 'invalid', 'PropertyTable': []}, {'results': [3], 'PropertyTable': {'Properties': [3]}}])
def test_malformed_government_responses_are_unavailable(monkeypatch, body):
    mock_async_client(monkeypatch, lambda request: httpx.Response(200, json=body))
    adapter = GlobalGovAPIClient()
    for result in (asyncio.run(adapter.fetch_openfda_drug_label('example')), asyncio.run(adapter.fetch_pubchem_compound_data('example'))):
        assert result['status'] == 'UNAVAILABLE' and result['records'] == []


def test_matching_quote_does_not_certify_interpretation(monkeypatch):
    monkeypatch.setenv('GROQ_API_KEY', 'test-only')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research, 'fetch_source', retrieve)
    claim = {'statement': 'Scope needs independent interpretation.', 'source_id': 'fda-mocra', 'quote': evidence(research.SOURCES[1])['text']}
    mock_async_client(monkeypatch, lambda request: httpx.Response(200, json={'choices': [{'message': {'content': json.dumps({'claims': [claim]})}}]}))
    result = asyncio.run(research.search('MoCRA', 'Cosmetics', 'FDA', 'en'))
    assert result['status'] == 'DRAFT'
    assert result['citation_match'] is True and result['interpretation_verified'] is False and result['verified'] is False


def test_irrelevant_topics_do_not_receive_generic_sources():
    assert research.select_sources('fictional medicine ZZZ-NONEXISTENT-9284 approval number', 'Pharmaceuticals', 'FDA') == []
    assert research.select_sources('Explain the procedure', 'Certification', 'EMA') == []


def test_qmsr_has_a_specific_current_official_source():
    sources = research.select_sources('QMSR and QSIT as of September 2026', 'Medical Devices', 'FDA')
    assert sources[0]['id'] == 'fda-qmsr'
    assert all(s['region'] == 'FDA' for s in sources)


def test_product_verdict_with_unrelated_exact_quote_is_rejected():
    claim = {'statement': 'This product is approved.', 'source_id': 'fda-mocra', 'quote': evidence(research.SOURCES[1])['text']}
    with pytest.raises(ValueError, match='Product verdict'):
        research.checked_claims(json.dumps({'claims': [claim]}), [evidence(research.SOURCES[1])])


def test_model_abstention_is_not_misreported_as_provider_failure(monkeypatch):
    monkeypatch.setenv('GROQ_API_KEY', 'test-only')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research, 'fetch_source', retrieve)
    mock_async_client(monkeypatch, lambda r: httpx.Response(200, json={'choices': [{'message': {'content': '{"claims": []}'}}]}))
    result = asyncio.run(research.search('MoCRA', 'Cosmetics', 'FDA', 'en'))
    assert result['status'] == 'INSUFFICIENT_EVIDENCE' and result['claims'] == []


def test_unchanged_foreign_language_translation_is_not_success(monkeypatch):
    monkeypatch.setenv('DEEPL_API_KEY', 'test-only')
    text = 'The sample contains 25 mg.'
    mock_async_client(monkeypatch, lambda r: httpx.Response(200, json={'translations': [{'text': text, 'detected_source_language': 'EN'}]}))
    with pytest.raises(translation.TranslationError, match='unchanged text'):
        asyncio.run(translation.translate(text, 'en', 'de'))


@pytest.mark.parametrize('name', ['VITE_GROQ_API_KEY', 'VITE_GROQ_API_KEY_1', 'VITE_GROQ_API_KEY_2', 'VITE_GROQ_API_KEY_3'])
def test_legacy_server_key_drives_real_provider_request_without_exposure(client, monkeypatch, name):
    monkeypatch.setenv(name, 'legacy-test-secret')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research, 'fetch_source', retrieve)
    def handler(request):
        assert request.headers['Authorization'] == 'Bearer legacy-test-secret'
        return httpx.Response(200, json={'choices': [{'message': {'content': '{"claims": []}'}}]})
    mock_async_client(monkeypatch, handler)
    result = asyncio.run(research.search('MoCRA', 'Cosmetics', 'FDA', 'en'))
    assert result['provider_status'] == 'RESPONDED'
    assert client.get('/api/health').json()['capabilities']['answer_synthesis'] == 'CONFIGURED_NOT_PROBED'
    assert 'legacy-test-secret' not in client.get('/').text
    assert 'legacy-test-secret' not in json.dumps(result)


def test_primary_key_precedes_legacy_and_no_quota_rotation(monkeypatch):
    from engine.ai_config import groq_key, configuration_status
    monkeypatch.setenv('GROQ_API_KEY', 'primary-test')
    monkeypatch.setenv('VITE_GROQ_API_KEY_1', 'legacy-test')
    assert groq_key() == 'primary-test'
    assert configuration_status() == {'configured': True, 'legacy_variable_in_use': False}


@pytest.mark.parametrize('code,reason', [(401, 'AUTHENTICATION_FAILED'), (403, 'ACCESS_DENIED'), (429, 'RATE_LIMITED'), (503, 'PROVIDER_ERROR')])
def test_ai_provider_failure_is_distinguishable_without_body_or_key_leak(monkeypatch, code, reason):
    monkeypatch.setenv('GROQ_API_KEY', 'private-test')
    async def retrieve(client, source): return evidence(source)
    monkeypatch.setattr(research, 'fetch_source', retrieve)
    mock_async_client(monkeypatch, lambda r: httpx.Response(code, json={'error': 'private-provider-details'}))
    result = asyncio.run(research.search('MoCRA', 'Cosmetics', 'FDA', 'en'))
    assert result['provider_status'] == reason and result['claims'] == []
    assert 'private-' not in json.dumps(result)
