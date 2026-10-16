"""Bounded in-memory document readers; no archive extraction or macro execution."""
import io
import re
import struct
import zlib
import zipfile
from pathlib import PurePosixPath
from defusedxml import ElementTree as ET
import olefile
from pypdf import PdfReader
from openpyxl import load_workbook
import xlrd
from striprtf.striprtf import rtf_to_text
from .translation_service import TranslationError

MAX_UPLOAD = 4_000_000
MAX_EXPANDED = 24_000_000
MAX_CHARS = 160_000
SUPPORTED = {'.pdf','.doc','.docx','.xlsx','.xls','.xlsm','.csv','.tsv','.hwp','.hwpx',
             '.pptx','.odt','.ods','.rtf','.txt','.md','.json','.xml','.html','.htm','.yaml','.yml','.log','.zip'}


def fail(message, status=422):
    raise TranslationError(message,status)


def bounded_zip(content):
    archive=zipfile.ZipFile(io.BytesIO(content))
    info=archive.infolist()
    if len(info)>1500 or sum(i.file_size for i in info)>MAX_EXPANDED:
        archive.close(); fail('Expanded archive exceeds the processing limit.',413)
    for item in info:
        name=item.filename.replace('\\','/')
        if item.flag_bits & 1 or name.startswith('/') or '..' in PurePosixPath(name).parts or (item.external_attr >> 16) & 0o170000 == 0o120000:
            archive.close(); fail('Encrypted archives, symbolic links and unsafe archive paths are not supported.')
        if item.file_size>MAX_EXPANDED or item.file_size>max(item.compress_size,1)*1000:
            archive.close(); fail('Archive compression or member size exceeds the safe limit.',413)
    return archive


def text_decode(content):
    if content.startswith((b'\xff\xfe',b'\xfe\xff')):
        return content.decode('utf-16')
    for encoding in ('utf-8-sig','cp949','cp1252'):
        try:
            value=content.decode(encoding)
            if '\x00' in value or sum(ord(c)<32 and c not in '\n\r\t' for c in value)>3:
                continue
            return value
        except UnicodeDecodeError:
            pass
    fail('This file contains unsupported binary data or an unreadable text encoding.')


def hwp_paragraph(data):
    # HWP v5 control codes use 1 or 8 UTF-16 code units, not printable payloads.
    out=[]; pos=0
    wide={1,2,3,4,5,6,7,8,9,11,12,14,15,16,17,18,19,20,21,22,23}
    while pos+2<=len(data):
        code=struct.unpack_from('<H',data,pos)[0]
        if code<32:
            if code in (10,13):out.append('\n')
            elif code==9:out.append('\t')
            pos+=16 if code in wide else 2
        else:
            start=pos
            while pos+2<=len(data) and struct.unpack_from('<H',data,pos)[0]>=32:pos+=2
            out.append(data[start:pos].decode('utf-16le',errors='strict'))
    return ''.join(out)


def hwp_sections(content):
    with olefile.OleFileIO(io.BytesIO(content)) as ole:
        header=ole.openstream('FileHeader').read()
        if not header.startswith(b'HWP Document File') or len(header)<40:fail('Only HWP version 5 documents are supported. Export older documents as HWPX or DOCX.')
        flags=struct.unpack_from('<I',header,36)[0]
        if flags & (2|4):fail('Password-protected or distribution-protected HWP files cannot be read.')
        paths=sorted((p for p in ole.listdir() if len(p)==2 and p[0]=='BodyText' and p[1].startswith('Section')),key=lambda p:int(p[1][7:]))
        total=0
        for index,path in enumerate(paths):
            data=ole.openstream(path).read(MAX_EXPANDED+1)
            if flags&1:
                dec=zlib.decompressobj(-15);data=dec.decompress(data,MAX_EXPANDED+1)
                if not dec.eof or dec.unconsumed_tail:fail('HWP expanded content exceeds the limit or is corrupt.',413)
            total+=len(data)
            if total>MAX_EXPANDED:fail('HWP expanded content exceeds the limit.',413)
            out=[];pos=0
            while pos<len(data):
                if pos+4>len(data):fail('Invalid HWP record.')
                header=struct.unpack_from('<I',data,pos)[0];pos+=4
                tag=header&1023;size=header>>20
                if size==4095:
                    size=struct.unpack_from('<I',data,pos)[0];pos+=4
                if pos+size>len(data):fail('Truncated HWP record.')
                if tag==67:out.append(hwp_paragraph(data[pos:pos+size]))
                pos+=size
            yield f'Section {index+1}', '\n'.join(out)


def doc_piece_text(word,table):
    flags=struct.unpack_from('<H',word,10)[0]
    if flags & 0x8100:fail('Encrypted DOC documents are not supported.')
    if struct.unpack_from('<H',word,2)[0]<0xc1:fail('DOC files older than Word 97 need conversion to DOCX.')
    pos=32;csw=struct.unpack_from('<H',word,pos)[0];pos+=2+csw*2
    cslw=struct.unpack_from('<H',word,pos)[0];pos+=2
    ccp_text=struct.unpack_from('<I',word,pos+12)[0]
    pos+=cslw*4
    count=struct.unpack_from('<H',word,pos)[0];pos+=2
    if count<=33:fail('DOC piece table is missing.')
    fc,lcb=struct.unpack_from('<II',word,pos+33*8)
    clx=table[fc:fc+lcb];i=0
    while i<len(clx) and clx[i]==1:i+=3+struct.unpack_from('<H',clx,i+1)[0]
    if i>=len(clx) or clx[i]!=2:fail('Unsupported DOC text layout. Save as DOCX.')
    length=struct.unpack_from('<I',clx,i+1)[0];plc=clx[i+5:i+5+length]
    if len(plc)!=length or (length-4)%12:fail('Corrupt DOC piece table.')
    n=(length-4)//12;out=[]
    for k in range(n):
        start,end=struct.unpack_from('<II',plc,k*4)
        if start>=ccp_text:break
        end=min(end,ccp_text)
        if end<start:fail('Invalid DOC character range.')
        fc=struct.unpack_from('<I',plc,(n+1)*4+k*8+2)[0]
        compressed=bool(fc&0x40000000);offset=fc&0x3fffffff
        if compressed:offset//=2
        size=(end-start)*(1 if compressed else 2)
        piece=word[offset:offset+size]
        if len(piece)!=size:fail('Truncated DOC text.')
        out.append(piece.decode('cp1252' if compressed else 'utf-16le',errors='strict'))
    text=''.join(out).replace('\r','\n').replace('\x07','\t')
    # Word fields: retain the displayed result, never evaluate field instructions.
    text=re.sub('\x13[^\x14\x15]*\x14','',text)
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',text)


def xml_blocks(content, mode):
    root=ET.fromstring(content)
    tags={'p'} if mode in ('docx','odt','ods') else {'p'}
    # Each leaf text node is included once; table paragraphs keep their order.
    for element in root.iter():
        local=element.tag.rsplit('}',1)[-1]
        if local in tags:
            if any(x is not element and x.tag.rsplit('}',1)[-1]=='p' for x in element.iter()):continue
            if mode=='docx':
                value=''.join((x.text or '') if x.tag.rsplit('}',1)[-1]=='t' else '\t' if x.tag.rsplit('}',1)[-1]=='tab' else '\n' if x.tag.rsplit('}',1)[-1]=='br' else '' for x in element.iter())
            else:value=''.join(element.itertext())
            if value.strip():yield value


def read_sections(name,content):
    ext=PurePosixPath(name.lower()).suffix
    notes=['Text translation only: original layout, images and embedded objects are not reproduced.']
    sections=[]
    if ext=='.pdf':
        reader=PdfReader(io.BytesIO(content))
        if reader.is_encrypted:fail('Password-protected PDF: remove protection before uploading.')
        if len(reader.pages)>150:fail('PDF exceeds 150 pages. Split it into smaller documents.',413)
        for i,page in enumerate(reader.pages):
            # Bound decompressed page streams before text processing.
            stream=page.get_contents()
            if stream is not None and len(stream.get_data())>MAX_EXPANDED:fail('PDF page content is too large.',413)
            text=page.extract_text() or ''
            if not text.strip():fail(f'PDF page {i+1} has no readable text. OCR is required; this document has not been partially translated.')
            sections.append((f'Page {i+1}',text))
        notes.append('Text may have PDF reading-order differences; verify tables, columns and symbols against the original.')
    elif ext in ('.docx','.hwpx','.pptx','.odt','.ods'):
        with bounded_zip(content) as z:
            names=z.namelist()
            if ext=='.docx':
                parts=['word/document.xml']+[p for p in names if re.fullmatch(r'word/(?:header\d+|footer\d+|footnotes|endnotes)\.xml',p)]
                for p in parts:sections.append((p,'\n'.join(xml_blocks(z.read(p),'docx'))))
                notes.append('Tracked revisions, comments and text in embedded objects require original-document review.')
            elif ext=='.hwpx':
                paths=sorted([p for p in names if re.fullmatch(r'Contents/section\d+\.xml',p)],key=lambda p:int(re.search(r'(\d+)\.xml',p)[1]))
                for p in paths:
                    root=ET.fromstring(z.read(p));text='\n'.join(''.join(n.itertext()) for n in root.iter() if n.tag.rsplit('}',1)[-1]=='t')
                    sections.append((p,text))
            elif ext=='.pptx':
                paths=sorted([p for p in names if re.fullmatch(r'ppt/(?:slides/slide|notesSlides/notesSlide)\d+\.xml',p)])
                for p in paths:
                    root=ET.fromstring(z.read(p));sections.append((p,'\n'.join(n.text or '' for n in root.iter() if n.tag.rsplit('}',1)[-1]=='t')))
            else:sections.append(('Document','\n'.join(xml_blocks(z.read('content.xml'),ext[1:]))))
    elif ext in ('.xlsx','.xlsm'):
        with bounded_zip(content):pass
        book=load_workbook(io.BytesIO(content),read_only=True,data_only=False,keep_links=False)
        try:
            for sheet in book:
                if sheet.max_row>20000 or sheet.max_column>200:fail('Spreadsheet exceeds 20,000 rows or 200 columns per sheet.',413)
                lines=[];char_count=0
                for row in sheet.iter_rows():
                    values=[f'{cell.coordinate}: {cell.value}' for cell in row if cell.value is not None]
                    if values:
                        line=' | '.join(values);lines.append(line);char_count+=len(line)
                    if char_count>MAX_CHARS:fail('Spreadsheet text is too large.',413)
                sections.append((sheet.title,'\n'.join(lines)))
        finally:book.close()
        notes.append('All worksheets, including hidden sheets, are read. Formula text is retained but never calculated. Macros are not executed.')
    elif ext=='.xls':
        book=xlrd.open_workbook(file_contents=content,on_demand=True)
        try:
            for sheet in book.sheets():
                if sheet.nrows*sheet.ncols>200000:fail('XLS worksheet is too large.',413)
                sections.append((sheet.name,'\n'.join(f'Row {i+1}: '+' | '.join(str(v) for v in sheet.row_values(i)) for i in range(sheet.nrows))))
        finally:book.release_resources()
        notes.append('XLS uses stored cell values; formulas are not recalculated. Dates may appear as stored numeric values.')
    elif ext=='.hwp':
        sections=list(hwp_sections(content));notes.append('HWP v5 body text and table paragraphs are extracted; images, comments and layout require review.')
    elif ext=='.doc':
        if content.lstrip().startswith(b'{\\rtf'):sections=[('Document',rtf_to_text(text_decode(content)))]
        else:
            with olefile.OleFileIO(io.BytesIO(content)) as ole:
                word=ole.openstream('WordDocument').read(MAX_EXPANDED+1)
                flags=struct.unpack_from('<H',word,10)[0]
                table=ole.openstream('1Table' if flags&0x200 else '0Table').read(MAX_EXPANDED+1)
                sections=[('Document body',doc_piece_text(word,table))]
            notes.append('Word 97–2003 main body text only; headers, footnotes, comments and embedded objects are excluded. Convert to DOCX to include more document parts.')
    elif ext=='.rtf':sections=[('Document',rtf_to_text(text_decode(content))) ]
    elif ext in SUPPORTED- {'.zip'}:
        value=text_decode(content)
        if ext in ('.html','.htm'):
            from bs4 import BeautifulSoup
            soup=BeautifulSoup(value,'html.parser')
            for n in soup(['script','style']):n.decompose()
            value=soup.get_text('\n',strip=True)
        sections=[('Document',value)]
    else:fail('This format can be attached, but text extraction is unavailable. Convert it to PDF (with text), DOCX, XLSX, HWPX or TXT. Images and scanned documents require OCR.',415)
    sections=[{'label':label,'text':text.strip()} for label,text in sections if text.strip()]
    if not sections:fail('No readable text was found. Scanned images require OCR.')
    if sum(len(s['text']) for s in sections)>MAX_CHARS:fail('Extracted document exceeds 160,000 characters. Split it before translating.',413)
    return {'name':name,'status':'EXTRACTED','sections':sections,'characters':sum(len(s['text']) for s in sections),'notes':notes}


def extract_upload(filename,content):
    if not content or len(content)>MAX_UPLOAD:fail('Each file must be non-empty and no larger than 4 MB.',413)
    name=(filename or 'attachment').replace('\\','/').split('/')[-1][:240]
    documents=[]
    def read_one(name,data):
        try:return read_sections(name,data)
        except TranslationError as e:return {'name':name,'status':'ERROR','message':str(e),'sections':[]}
        except Exception:return {'name':name,'status':'ERROR','message':'The file is damaged, encrypted or uses an unsupported variant. Export it as DOCX, XLSX, HWPX or a text PDF.','sections':[]}
    try:
        if name.lower().endswith('.zip'):
            with bounded_zip(content) as z:
                members=[i for i in z.infolist() if not i.is_dir() and not i.filename.startswith('__MACOSX/')]
                if len(members)>30:fail('A ZIP can contain at most 30 files.',413)
                total=0
                for item in members:
                    if item.filename.lower().endswith('.zip'):
                        documents.append({'name':item.filename,'status':'ERROR','message':'Nested ZIP files must be unpacked and uploaded separately.','sections':[]});continue
                    if item.file_size>MAX_UPLOAD:
                        documents.append({'name':item.filename,'status':'ERROR','message':'Archive member exceeds 4 MB.','sections':[]});continue
                    doc=read_one(item.filename,z.read(item))
                    total+=doc.get('characters',0)
                    if total>MAX_CHARS:fail('Combined archive text exceeds 160,000 characters. Split the archive.',413)
                    documents.append(doc)
        else:documents.append(read_one(name,content))
    except TranslationError:raise
    except Exception:fail('The archive could not be read.')
    if not documents:fail('The archive contains no documents.')
    return {'documents':documents,'status':'READY' if all(d['status']=='EXTRACTED' for d in documents) else 'PARTIAL_OR_UNSUPPORTED',
            'storage':'request_lifetime_only','original_files_retained':False}
