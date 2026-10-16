import asyncio
import io
import json
import struct
import zipfile
from unittest.mock import patch
import pytest
import httpx
from fastapi.testclient import TestClient
from docx import Document
from openpyxl import Workbook,load_workbook
from pypdf import PdfWriter
import app
from certification import file_reader as reader, file_translation as translator
from certification.translation_service import TranslationError


def archive(files):
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        for name,data in files.items():z.writestr(name,data)
    return b.getvalue()

def test_docx_body_table_and_header_are_extracted():
    d=Document();d.add_paragraph('Good Manufacturing Practice');d.add_table(rows=1,cols=1).cell(0,0).text='CAPA';d.sections[0].header.paragraphs[0].text='Quality system'
    b=io.BytesIO();d.save(b)
    result=reader.extract_upload('regulation.docx',b.getvalue())['documents'][0]
    text=' '.join(s['text'] for s in result['sections'])
    assert all(t in text for t in ['Good Manufacturing Practice','CAPA','Quality system'])

def test_xlsx_all_sheets_formulas_not_executed():
    book=Workbook();book.active['A1']='GMP';book.active['B1']='=1+2';sheet=book.create_sheet('Hidden');sheet.sheet_state='hidden';sheet['A1']='Hidden requirements';b=io.BytesIO();book.save(b)
    data=reader.extract_upload('terms.xlsx',b.getvalue())['documents'][0]
    assert len(data['sections'])==2
    assert '=1+2' in data['sections'][0]['text']
    assert 'Hidden requirements' in data['sections'][1]['text']

def test_hwpx_sections_in_numeric_order():
    xml=lambda s:f'<root xmlns:h="urn:test"><h:p><h:t>{s}</h:t></h:p></root>'
    data=reader.extract_upload('x.hwpx',archive({'Contents/section10.xml':xml('ten'),'Contents/section2.xml':xml('two')}))
    assert [s['text'] for s in data['documents'][0]['sections']]==['two','ten']

def test_hwp_record_reader_uses_body_not_preview():
    text='제조 및 품질관리 GMP'.encode('utf-16le');record=struct.pack('<I',67|(len(text)<<20))+text
    header=b'HWP Document File'+b'\0'*15+b'\x00\x00\x00\x05'+struct.pack('<I',0)
    class Ole:
        def __enter__(self):return self
        def __exit__(self,*args):pass
        def listdir(self):return [['BodyText','Section0']]
        def openstream(self,path):return io.BytesIO(header if path=='FileHeader' else record)
    with patch.object(reader.olefile,'OleFileIO',return_value=Ole()):
        assert list(reader.hwp_sections(b'fake'))[0][1]=='제조 및 품질관리 GMP'

def test_hwp_controls_are_not_rendered_as_text():
    data='GMP'.encode('utf-16le')+struct.pack('<H',9)+b'\0'*14+'CAPA'.encode('utf-16le')
    assert reader.hwp_paragraph(data)=='GMP\tCAPA'

def test_doc_piece_table():
    word=bytearray(1024);struct.pack_into('<H',word,2,0xc1);struct.pack_into('<H',word,32,0)
    struct.pack_into('<H',word,34,4);struct.pack_into('<I',word,48,3);struct.pack_into('<H',word,52,34)
    piece=struct.pack('<II',0,3)+b'\0\0'+struct.pack('<I',0x40000000|1200)+b'\0\0'
    clx=b'\x02'+struct.pack('<I',len(piece))+piece;struct.pack_into('<II',word,54+33*8,0,len(clx));word[600:603]=b'GMP'
    assert reader.doc_piece_text(bytes(word),clx)=='GMP'

def test_zip_handles_unsupported_members_explicitly():
    data=reader.extract_upload('files.zip',archive({'one.txt':'GMP requirements','unsupported.exe':b'\0\1'}))
    assert data['status']=='PARTIAL_OR_UNSUPPORTED'
    assert [d['status'] for d in data['documents']]==['EXTRACTED','ERROR']

def test_zip_never_extracts_unsafe_paths():
    with pytest.raises(TranslationError):reader.extract_upload('files.zip',archive({'../secret.txt':'test'}))

def test_nested_zip_is_reported():
    data=reader.extract_upload('files.zip',archive({'nested.zip':archive({'a.txt':'test'})}))
    assert data['documents'][0]['status']=='ERROR'

def test_scanned_pdf_is_not_silently_partial():
    w=PdfWriter();w.add_blank_page(width=100,height=100);b=io.BytesIO();w.write(b)
    data=reader.extract_upload('scan.pdf',b.getvalue())
    assert data['documents'][0]['status']=='ERROR'
    assert 'OCR' in data['documents'][0]['message']

def test_rtf_and_korean_plain_text():
    assert 'GMP' in reader.extract_upload('file.rtf',b'{\\rtf1\\ansi GMP}')['documents'][0]['sections'][0]['text']
    assert reader.extract_upload('file.txt','품질관리'.encode('cp949'))['documents'][0]['sections'][0]['text']=='품질관리'

def test_upload_limit():
    with pytest.raises(TranslationError):reader.extract_upload('big.txt',b'x'*(reader.MAX_UPLOAD+1))

def mock_provider(monkeypatch,translation='Maintain 5 mg of the substance.',finish='stop'):
    monkeypatch.setenv('GROQ_API_KEY','test-only')
    original=httpx.AsyncClient
    def handler(request):
        body=json.loads(request.content)
        assert 'English base term' in body['messages'][0]['content']
        return httpx.Response(200,json={'choices':[{'finish_reason':finish,'message':{'content':json.dumps({'translation':translation,'terms':[],'notes':[]})}}]})
    monkeypatch.setattr(httpx,'AsyncClient',lambda **kw:original(transport=httpx.MockTransport(handler),**kw))

def test_translation_preserves_numbers(monkeypatch):
    mock_provider(monkeypatch)
    result=asyncio.run(translator.translate_chunk(translator.ChunkPayload(text='Maintain 5 mg.',consent=True)))
    assert result['numeric_check']=='passed'
    assert result['certified'] is False

def test_missing_numeric_value_rejects_translation(monkeypatch):
    mock_provider(monkeypatch,'Maintain the substance.')
    with pytest.raises(TranslationError,match='number'):asyncio.run(translator.translate_chunk(translator.ChunkPayload(text='Maintain 5 mg.',consent=True)))

def test_truncated_provider_response_is_not_success(monkeypatch):
    mock_provider(monkeypatch,finish='length')
    with pytest.raises(TranslationError,match='incomplete'):asyncio.run(translator.translate_chunk(translator.ChunkPayload(text='GMP',consent=True)))

def test_consent_is_required():
    with pytest.raises(TranslationError):asyncio.run(translator.translate_chunk(translator.ChunkPayload(text='GMP')))

def test_office_export_no_formula_execution():
    p=translator.ExportPayload(format='xlsx',target_lang='ko',segments=[translator.ExportSegment(file='=HYPERLINK("evil")',location='A1',original='=1+2',translation='GMP')])
    content,mime=translator.export_translation(p)
    book=load_workbook(io.BytesIO(content));assert book.active['A4'].data_type=='s';assert book.active['C4'].data_type=='s'
    assert 'PARTIAL' in book.active['A1'].value
    p.format='docx';content,_=translator.export_translation(p);assert Document(io.BytesIO(content)).paragraphs[0].text.endswith('draft')

def test_api_extract_allows_files_without_question():
    app.LIMITS.clear()
    with TestClient(app.app,base_url='https://testserver') as client:
        r=client.post('/api/files/extract',files={'file':('terms.txt',b'Good Manufacturing Practice')})
        assert r.status_code==200
        assert r.json()['documents'][0]['status']=='EXTRACTED'
        assert client.post('/api/files/translate',json={'text':'GMP','consent':False}).status_code==400


def test_protected_units_and_identifiers():
    assert translator.protected_tokens('C95931 uses 5 mg and 200 µL.') == translator.protected_tokens('C95931: 5mg, 200µL.')
    assert translator.protected_tokens('C95931 5 mg') - translator.protected_tokens('C95932 5 g')
